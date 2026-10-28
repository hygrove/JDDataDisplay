# -*- coding: utf-8 -*-
"""库存表 ETL：把 ResourceData/库存数据表/*.xlsx 按列名白名单投影进 inventory_daily。

设计要点（贴合项目既有风格）：
- 列名白名单来自 inventory_metrics（唯一事实来源），ETL 与建表共用，绝不手写 182 列；
- 城市精确等值匹配：8 城白名单，衍生前缀「广州补货B」「上海B仓(商)京东自营渠道」自然不命中；
- 数值列缺失/空值写 0；文本列（身份列）缺失/空值写 NULL；
- 医拜缺字段（16-30/151-180 系列）列不存在 → 写 0；
- SKU 浮点清洗："100149827444.0" → "100149827444"；
- shop 由文件名前缀解析（家配/医拜/卡求），不读表内字段；
- 增量：文件级 sha256 门控，哈希变化才重算该文件对应的 (shop, date) 行。

为什么不用 polars：库存表是「每文件一张宽快照」，openpyxl 直读 + 列名投影更直观，
也避免把 2000+ 列宽表喂给 polars 的长表逻辑。
"""
from __future__ import annotations

import hashlib
import re
import sqlite3
from datetime import datetime, timezone
from pathlib import Path

from . import config
from .inventory_metrics import (
    all_columns,
    CITIES,
    CITY_CODES,
    NUMERIC_FIELDS,
    extract_city_slow_cols,
    source_to_field,
)


def _sha256(path: Path) -> str:
    """计算文件 sha256 十六进制串（按块读取）。"""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 16), b""):
            h.update(chunk)
    return h.hexdigest()


def _parse_shop(filename: str) -> str:
    """从文件名解析店铺名：「2026-10-18家配自营库存明细.xlsx」→「家配」。

    Args:
        filename: 文件名（含扩展名）。

    Returns:
        str: 店铺名；解析失败兜底返回去日期后的文件名。
    """
    stem = Path(filename).stem
    m = re.match(r"^\d{4}-\d{2}-\d{2}(.+?)自营库存明细", stem)
    if m:
        return m.group(1)
    # 兜底：去掉开头的 YYYY-MM-DD
    rest = re.sub(r"^\d{4}-\d{2}-\d{2}", "", stem)
    return rest or stem


def _clean_sku(value) -> str:
    """清洗 SKU：浮点串 "100149827444.0" → "100149827444"；其余原样转字符串。"""
    if value is None:
        return ""
    if isinstance(value, float):
        return str(int(value))
    s = str(value).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s


def _to_num(value):
    """把任意单元格值规整成 float；None/空串/非数字 → 0.0（数值列缺值写 0）。"""
    if value is None:
        return 0.0
    if isinstance(value, str):
        s = value.strip()
        if s == "":
            return 0.0
        try:
            return float(s)
        except ValueError:
            return 0.0
    if isinstance(value, bool):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    return 0.0


def _to_text(value):
    """文本列规整：None/空串 → None（NULL）；否则字符串。

    整数值的浮点（如医拜表 SPU 列 100149827444.0）去掉小数尾巴，
    避免页面上出现「xxx.0」这种脏展示。
    """
    if value is None:
        return None
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    if isinstance(value, str):
        s = value.strip()
        return s or None
    return str(value)


def _parse_date(value) -> str:
    """把「时间」列值规整成 YYYY-MM-DD 字符串。"""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    s = str(value).strip()
    # 形如 2026-10-18 00:00:00 → 取前 10 位
    if len(s) >= 10 and s[4] == "-" and s[7] == "-":
        return s[:10]
    return s


def _parse_file(path: Path) -> tuple[str, str, list[dict], list[dict]] | None:
    """解析单个库存 xlsx：投影成 (shop, date, 行字典列表, 全城市滞销长表行)。

    除 8 城白名单宽表外，还把源表里**全部**「<城市>滞销件数」列（约 70 个
    城市/仓）抽成 inventory_city_slow 长表行（只存 >0 的，控制体量）。

    跳过规则：锁文件、缺「时间」列、无数据行的空模板。

    Args:
        path: xlsx 文件路径。

    Returns:
        tuple[str, str, list[dict], list[dict]] | None: 解析成功返回
            (店铺, 日期, 宽表行列表, 滞销长表行列表)；应跳过时返回 None。
    """
    import openpyxl

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(min_row=1, values_only=True))
        wb.close()
    except Exception:
        # 单个文件损坏不拖垮整批，跳过
        return None
    if not rows:
        return None

    header = list(rows[0])
    if "时间" not in header:
        return None  # 非库存明细（如 POP 模板）
    idx = {h: i for i, h in enumerate(header) if h is not None}
    # 全城市滞销列（动态提取，不限 8 城白名单）
    slow_cols = extract_city_slow_cols(header)

    src_map = source_to_field()
    numeric_set = set(NUMERIC_FIELDS)
    date_value = None
    shop_value = _parse_shop(path.name)
    out: list[dict] = []
    slow_out: list[dict] = []
    for raw in rows[1:]:
        if raw is None or all(v is None for v in raw):
            continue  # 空行
        rec: dict = {}
        for fld, src, ctype in all_columns():
            if src is None:
                continue  # shop / erp_code 稍后填
            col = idx.get(src)
            if col is None:
                # 列不存在（如医拜缺 16-30 系列）→ 数值写 0，文本写 NULL
                rec[fld] = 0.0 if ctype == "real" else None
            else:
                v = raw[col] if col < len(raw) else None
                if ctype == "real":
                    rec[fld] = _to_num(v)
                else:
                    rec[fld] = _to_text(v)
        # 主键与身份字段
        rec["sku"] = _clean_sku(rec.get("sku"))
        if not rec["sku"]:
            continue
        rec["date"] = date_value if date_value else _parse_date(raw[idx["时间"]] if idx.get("时间") is not None and idx["时间"] < len(raw) else None)
        if date_value is None:
            date_value = rec["date"]
        rec["shop"] = shop_value
        rec["erp_code"] = None
        out.append(rec)

        # 全城市滞销长表：仅 >0 的 (城市, 件数) 入表
        for city, ci in slow_cols:
            v = _to_num(raw[ci] if ci < len(raw) else None)
            if v > 0:
                slow_out.append({
                    "date": rec["date"], "shop": shop_value,
                    "sku": rec["sku"], "city": city, "slow_units": v,
                })

    if not out or not date_value:
        return None
    return out[0]["shop"], date_value, out, slow_out


def _replace_file_rows(
    shop: str, date: str, rows: list[dict], slow_rows: list[dict], key: str, h: str
) -> None:
    """删除该文件对应 (shop, date) 的旧行，重插新行，并刷新哈希。

    同时维护 inventory_city_slow（全城市滞销长表）：先删后插，与宽表同事务。
    """
    from .db import get_conn

    cols = [c[0] for c in all_columns()]
    col_sql = ", ".join(f'"{c}"' for c in cols)
    placeholders = ", ".join(["?"] * len(cols))
    data = [[r.get(c) for c in cols] for r in rows]
    slow_data = [[r["date"], r["shop"], r["sku"], r["city"], r["slow_units"]] for r in slow_rows]
    conn = get_conn()
    try:
        conn.execute("DELETE FROM inventory_daily WHERE shop=? AND date=?", (shop, date))
        conn.executemany(
            f'INSERT OR REPLACE INTO inventory_daily ({col_sql}) VALUES ({placeholders})', data
        )
        conn.execute(
            "DELETE FROM inventory_city_slow WHERE shop=? AND date=?", (shop, date)
        )
        conn.executemany(
            "INSERT OR REPLACE INTO inventory_city_slow(date, shop, sku, city, slow_units) VALUES (?,?,?,?,?)",
            slow_data,
        )
        now = datetime.now(timezone.utc).isoformat()
        conn.execute(
            "INSERT OR REPLACE INTO inventory_processed_file(rel_path, content_hash, processed_at) VALUES (?,?,?)",
            (key, h, now),
        )
        conn.commit()
    finally:
        conn.close()


def ingest_inventory_to_db(root: Path | None = None,
                           only_files: set[str] | None = None) -> dict[str, object]:
    """把库存 xlsx 增量摄入 SQLite（文件级哈希门控）。

    Args:
        root: 库存数据源目录；默认 config.INVENTORY_SOURCE_DIR。
        only_files: 只摄入这些**文件名**（集合）；None = 摄入目录里全部文件。
            用途（ADR-0013）：库存按「每店保留最新一份」淘汰旧快照时，
            只该摄入保留集内的文件——省掉对"马上就要被删的旧文件"的全量解析
            （家配单份 1.66MB / 2000+ 列宽表，每份解析代价不小）。

    Returns:
        dict[str, object]: {changed, files, rows, changed_files, reason,
            shop_rows}。`shop_rows` = {店铺: 本次写入行数}，供健康检查门控使用
            （门控要判「新快照行数是否够」，不能只看"文件在不在"）。
    """
    from .db import ensure_inventory_schema, get_conn

    root = root or config.INVENTORY_SOURCE_DIR
    ensure_inventory_schema()
    files = sorted(p for p in Path(root).glob("*.xlsx") if not p.name.startswith("~$"))
    if only_files is not None:
        files = [p for p in files if p.name in only_files]
    if not files:
        return {"changed": False, "files": 0, "rows": 0, "changed_files": 0,
                "shop_rows": {}, "reason": "no_source_files"}

    hashes = {p.name: _sha256(p) for p in files}
    conn = get_conn()
    try:
        cur = conn.execute("SELECT rel_path, content_hash FROM inventory_processed_file")
        old = {r["rel_path"]: r["content_hash"] for r in cur.fetchall()}
        has_rows = conn.execute("SELECT COUNT(*) AS c FROM inventory_daily").fetchone()["c"]
    finally:
        conn.close()

    total_rows = 0
    changed_files = 0
    shop_rows: dict[str, int] = {}   # 本次各店写入行数（供健康检查门控用）
    for p in files:
        key = p.name
        if key in old and hashes[key] == old[key] and has_rows > 0:
            continue  # 哈希一致且库非空 → 跳过（增量收益点）
        parsed = _parse_file(p)
        if parsed is None:
            continue
        shop, date, rows, slow_rows = parsed
        _replace_file_rows(shop, date, rows, slow_rows, key, hashes[key])
        total_rows += len(rows)
        changed_files += 1
        # 该店这一份的入库行数（门控要判「新快照行数是否够」，
        # 用本次实际写入的行数比用"文件在不在"准确）
        shop_rows[shop] = shop_rows.get(shop, 0) + len(rows)

    return {
        "changed": changed_files > 0,
        "files": len(files),
        "rows": total_rows,
        "changed_files": changed_files,
        "shop_rows": shop_rows,
        "reason": "ingested" if changed_files else "unchanged",
    }
