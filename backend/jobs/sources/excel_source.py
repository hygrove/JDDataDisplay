# -*- coding: utf-8 -*-
"""POP 单品明细数据源：读固定目录下的 Excel 明细 + CSV 推广，产出统一长表。

目录约定（项目内固定路径，不再按「最新日期文件夹」遍历）：
  {root}/数据源表目录/{店铺名}_商品明细_{起}_{止}.xlsx
  {root}/数据源表目录/{店铺名}_推广数据_{起}_{止}.csv
其中 {root} = config.POP_SOURCE_DIR（默认项目下 ResourceData）。明细为 xlsx，推广为 csv。

fetch() 主流程：
  1. 读所有店铺的 xlsx 明细，纵向拼接成长表；
  2. 按 (shop, date, spu) **左连接** csv 推广数据，补充 推广花费 / 推广成交金额；
  3. 补齐 LONG_COLUMNS 缺失列、统一列顺序；
  4. 清洗①去重（同一键保留最后出现的）；
  5. 清洗②白名单过滤（只保留「店铺spu登记信息.xlsx」里登记过的 SPU）；
  6. 用「SPU商品名称映射表」的简化名覆盖原名称（映射表没有则保留原名）。

⚠️ 明细与推广两份数据必须做**完全一致**的去重 + 白名单过滤，
  否则 JOIN 后会出现「有推广费但没明细」或反之的错行。
"""
from __future__ import annotations

from pathlib import Path

import polars as pl

from .. import config
from .base import LONG_COLUMNS

# COLUMN_MAP：源表中文列名 -> 统一长表英文列名。
# 只映射「可累加」的原始指标；比率类不进长表。
COLUMN_MAP = {
    "时间": "date",
    "SPU": "spu",
    "SPU名称": "spu_name",
    "三级类目": "category",
    "商品访客数": "visitors",
    "成交客户数": "buyers",
    "成交单量": "orders",
    "成交商品件数": "items",
    "成交金额": "amount",
    "搜索曝光次数": "search_impressions",
    "搜索点击次数": "search_clicks",
    "取消及售后退款单量": "refund_orders",
    "取消及售后退款金额": "refund_amount",
    # 说明：「成交转化率」「客单价」「搜索点击率」为比率类派生指标，源表虽有同名列，
    # 但汇总时不能直接相加/取平均，统一由 transforms.compute_metrics 重算，故不入长表。
}

# NAME_MAP_FILE：商品名称简化映射表，位于数据源**根目录**（不是日期子目录）。
# 用途：源表的 SPU 名称往往冗长带营销词，运营维护一张简名表让页面更好读
NAME_MAP_FILE = "SPU商品名称映射表.xlsx"

# SPU_REGISTER_FILE：店铺 SPU 登记信息（数据源根目录）。
# 结构：第一行店铺名作列头，其下每列是该店铺需要的 SPU。
# ⚠️ 这实际上是一道**白名单闸门**——不在表里的 SPU 会被整条链路过滤掉，
#    所以「改了数据源但页面没出现新店铺/新 SPU」时，第一件事是检查这里有没有登记
SPU_REGISTER_FILE = "店铺spu登记信息.xlsx"


def load_spu_name_map(root: Path | None = None) -> dict[str, str]:
    """读取「SPU商品名称映射表.xlsx」，返回 {spu 编号(str): 简化商品名称}。

    映射表结构：Sheet1 含 `spu` 与 `商品名称` 两列（表头必须精确匹配这两个名字）。

    Args:
        root: 数据源根目录；可选，默认取 config.POP_SOURCE_DIR。

    Returns:
        dict[str, str]: SPU 编号 -> 简化名称。
            文件不存在、表头不符、读取异常时一律返回空 dict（降级使用明细表原名称），
            **不抛异常**——映射表是可选增强，不应中断批处理。

    Example:
        >>> load_spu_name_map(Path("ResourceData"))  # doctest: +SKIP
        {'100123': '钻芯保温杯'}
    """
    src = (root or config.POP_SOURCE_DIR) / NAME_MAP_FILE
    if not src.is_file():
        return {}
    try:
        # openpyxl 延迟导入：只有真的要读映射表时才付出导入成本
        import openpyxl

        # read_only + data_only：只读值不读公式，避免拿到公式字符串
        wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return {}
        header = [str(c).strip() if c is not None else "" for c in rows[0]]
        # 表头校验：两列都必须存在，否则宁可返回空也不要瞎猜列位置
        if "spu" not in header or "商品名称" not in header:
            return {}
        si, ni = header.index("spu"), header.index("商品名称")
        mapping: dict[str, str] = {}
        for r in rows[1:]:
            # 跳过空行与任一关键列为空的行
            if not r or r[si] is None or r[ni] is None:
                continue
            key = str(r[si]).strip()
            val = str(r[ni]).strip()
            if key:
                mapping[key] = val
        return mapping
    except Exception:
        # 映射表是可选增强，任何异常都不应中断批处理
        return {}


def _norm_spu(v: object) -> str:
    """把 SPU 值归一成干净字符串。

    为什么要归一：
      Excel 里 SPU 常被存成数字，读出来变成 "100123.0" 或科学计数法 "1.00123E5"，
      而推广 CSV 里是纯字符串 "100123"。若不归一，两边 JOIN 会全部失配（推广数据全空）。

    Args:
        v: 任意来源的 SPU 值（int / float / str / None）。

    Returns:
        str: 归一后的 SPU 字符串；整数值去掉小数尾巴，非数字原样返回。

    Example:
        >>> _norm_spu(100123.0)
        '100123'
        >>> _norm_spu(" 100123 ")
        '100123'
    """
    s = str(v).strip()
    if s.endswith(".0"):
        s = s[:-2]
    try:
        f = float(s)
        if f.is_integer():
            return str(int(f))
    except (ValueError, OverflowError):
        # 非数字（如带字母的编码）保持原样
        pass
    return s


def load_spu_whitelist(root: Path | None = None) -> dict[str, set[str]]:
    """读取「店铺spu登记信息.xlsx」，返回 {店铺名: {spu, ...}} 白名单。

    结构：第一行是店铺名称（作列头），其下每列是该店铺需要的 SPU，逐行往下登记。

    Args:
        root: 数据源根目录；可选，默认取 config.POP_SOURCE_DIR。

    Returns:
        dict[str, set[str]]: 店铺名 -> 该店登记的 SPU 集合。
            文件不存在或读取异常时返回空 dict，表示**不过滤**（保留全部 SPU）。

    Example:
        >>> load_spu_whitelist(Path("ResourceData"))  # doctest: +SKIP
        {'钻芯旗舰店': {'100123', '100456'}}
    """
    src = (root or config.POP_SOURCE_DIR) / SPU_REGISTER_FILE
    if not src.is_file():
        return {}
    try:
        import openpyxl

        wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return {}
        header = [str(c).strip() if c is not None else "" for c in rows[0]]
        result: dict[str, set[str]] = {}
        for ri in range(1, len(rows)):
            row = rows[ri]
            for ci, shop in enumerate(header):
                # 空列头（Excel 尾部多余列）直接跳过，避免生成空店铺键
                if not shop:
                    continue
                # 行长度可能短于表头（尾部单元格为空），越界时按 None 处理
                val = row[ci] if ci < len(row) else None
                if val is None:
                    continue
                spu = _norm_spu(val)
                if spu:
                    result.setdefault(shop, set()).add(spu)
        return result
    except Exception:
        # 登记信息为可选裁剪，任何异常都不应中断批处理
        return {}


class PopSpuExcelSource:
    """POP 单品明细数据源：读取固定「数据源表目录」下各店铺的 Excel 明细 + CSV 推广。

    Attributes:
        name: 数据源标识，固定为 "pop_spu_excel"。
        root: 数据源根目录（Path），默认 config.POP_SOURCE_DIR。
        PROMO_COL_MAP: 推广 CSV 的中文列名 -> 长表列名映射（类常量，见下方定义）。

    Example:
        >>> src = PopSpuExcelSource()                    # doctest: +SKIP
        >>> df = src.fetch()                             # doctest: +SKIP
        >>> df.columns[:3]                               # doctest: +SKIP
        ['shop', 'date', 'spu']
    """

    name = "pop_spu_excel"

    def __init__(self, root: Path | None = None):
        """初始化数据源。

        Args:
            root: 数据源根目录；可选，默认 config.POP_SOURCE_DIR。
                测试时可传入临时目录以隔离真实数据。

        Returns:
            None
        """
        self.root = root or config.POP_SOURCE_DIR

    def fetch(self) -> pl.DataFrame:
        """拉取全部店铺明细并关联推广数据，返回清洗后的统一长表。

        Returns:
            pl.DataFrame: 列严格等于 LONG_COLUMNS 的长表，已完成去重与白名单过滤。

        Raises:
            FileNotFoundError: 数据源表目录不存在，或目录下没有任何 .xlsx 文件时抛出
                （此时继续跑下去只会产出空数据，不如显式报错让人检查数据源）。

        Example:
            >>> df = PopSpuExcelSource(Path("ResourceData")).fetch()   # doctest: +SKIP
            >>> df.height > 0                                          # doctest: +SKIP
            True
        """
        # 固定文件夹：{root}/数据源表目录（不再遍历最新日期文件夹，保证稳定性）
        table_dir = self.root / config.POP_TABLE_DIR_NAME
        if not table_dir.is_dir():
            raise FileNotFoundError(f"数据源表目录不存在：{table_dir}")
        frames: list[pl.DataFrame] = []
        # 跳过 Excel 打开时生成的 ~$ 临时锁文件（不是有效工作簿，读它会抛 CalamineError）
        for path in sorted(p for p in table_dir.glob("*.xlsx") if not p.name.startswith("~$")):
            # 文件名约定「{店铺名}_商品明细_...」，取第一段作为店铺名
            shop = path.name.split("_")[0]
            frames.append(self._read_one(path, shop))
        if not frames:
            raise FileNotFoundError(f"{table_dir} 下没有 xlsx 文件")
        df = pl.concat(frames, how="diagonal_relaxed")

        # —— 关联推广数据 CSV —— 按 (shop, date, spu) 左连接，补充 推广费/推广成交金额
        promo = self._fetch_promotion()
        if promo is not None:
            # 先丢弃明细里可能已存在的同名列，避免 JOIN 时报「列重复」
            drop_cols = [c for c in ("promotion_cost", "promotion_amount") if c in df.columns]
            if drop_cols:
                df = df.drop(drop_cols)
            df = df.join(promo, on=["shop", "date", "spu"], how="left")

        # 补齐缺失列，统一列顺序
        for col in LONG_COLUMNS:
            if col not in df.columns:
                df = df.with_columns(pl.lit(None).alias(col))
        df = df.select(LONG_COLUMNS)

        # —— 清洗 1：去重 —— 多份明细表日期区间重叠会产生 (店铺, 日期, SPU) 重复行，
        # 以这三列作为唯一键，保留最后一次出现（文件名排序更靠后 / 日期区间更新的那份）。
        df = df.unique(subset=["shop", "date", "spu"], keep="last")

        # —— 清洗 2：多余数据过滤 —— 只保留「店铺spu登记信息.xlsx」里登记的 SPU
        whitelist = load_spu_whitelist(self.root)
        if whitelist:
            # 预先把「店铺+SPU」二元组展开成集合，避免逐行重复做字符串归一
            allowed = {
                (shop.strip(), _norm_spu(spu))
                for shop, spus in whitelist.items()
                for spu in spus
            }
            df = df.filter(
                pl.struct(["shop", "spu"]).map_elements(
                    lambda s: (s["shop"].strip(), _norm_spu(s["spu"])) in allowed,
                    return_dtype=pl.Boolean,
                )
            )

        # 优先使用名称映射表的简化名；映射表没有时保留明细表原名称
        name_map = load_spu_name_map(self.root)
        if name_map:
            map_df = pl.DataFrame(
                {"spu": list(name_map.keys()), "spu_name_map": list(name_map.values())},
                schema={"spu": pl.String, "spu_name_map": pl.String},
            )
            # 左连接后用 coalesce 语义：映射名存在则用映射名，否则回退原名
            df = df.join(map_df, on="spu", how="left").with_columns(
                pl.when(pl.col("spu_name_map").is_not_null())
                .then(pl.col("spu_name_map"))
                .otherwise(pl.col("spu_name"))
                .alias("spu_name")
            ).drop("spu_name_map")
        return df

    def _read_one(self, path: Path, shop: str) -> pl.DataFrame:
        """读取单个店铺的明细 xlsx，规整成部分长表（不含推广列）。

        Args:
            path: xlsx 文件路径。
            shop: 该文件的店铺名（由文件名推断后传入），会作为 shop 列写入。

        Returns:
            pl.DataFrame: 已重命名、过滤、类型转换后的明细表。

        Raises:
            Exception: 文件损坏或不是合法工作簿时，由 polars.read_excel 抛出
                （上层 fetch 不捕获，让问题显式暴露）。

        Example:
            >>> df = src._read_one(Path("钻芯旗舰店_商品明细.xlsx"), "钻芯旗舰店")  # doctest: +SKIP
        """
        raw = pl.read_excel(path, sheet_id=0)
        # read_excel(sheet_id=0) 返回 {sheet_name: DataFrame}，取第一个 sheet
        df = next(iter(raw.values())) if isinstance(raw, dict) else raw

        # 只保留 COLUMN_MAP 里登记过的列（源表常带一堆无关列），并重命名为统一名
        keep = [c for c in COLUMN_MAP if c in df.columns]
        df = df.select(keep).rename(COLUMN_MAP)

        # 剔除合计行 / 空行，只保留单日记录：
        # 导出的表尾部常有「合计」汇总行，且日期列可能是别的格式，用正则锚定 YYYY-MM-DD
        df = df.filter(
            pl.col("spu").is_not_null()
            & (pl.col("spu").cast(pl.String) != "合计")
            & pl.col("date").cast(pl.String).str.contains(r"^\d{4}-\d{2}-\d{2}$")
        )
        df = df.with_columns(
            pl.lit(shop).alias("shop"),
            # SPU 必须归一：见 _norm_spu 说明，否则与推广表 JOIN 会失配
            pl.col("spu").cast(pl.String).map_elements(_norm_spu, return_dtype=pl.String),
            pl.col("date").cast(pl.String),
            pl.col("spu_name").cast(pl.String),
            pl.col("category").cast(pl.String),
        )
        # 指标列转数值（源表可能混入字符串/百分比/千分位，如 "1,234" 或 "12%"）
        for col in ("visitors", "buyers", "orders", "items", "amount",
                    "search_impressions", "search_clicks",
                    "refund_orders", "refund_amount"):
            if col in df.columns:
                df = df.with_columns(
                    pl.col(col).cast(pl.String)
                    .str.replace_all(",", "")   # 去千分位
                    .str.replace_all("%", "")   # 去百分号
                    # strict=False：无法解析时给 null 而不是抛异常，脏数据不拖垮整表
                    .cast(pl.Float64, strict=False)
                    .alias(col)
                )
        # 退款为空表示当天没有退款（JD 导出留空），按 0 处理，避免表格里大面积显示 --
        for col in ("refund_orders", "refund_amount"):
            if col in df.columns:
                df = df.with_columns(pl.col(col).fill_null(0).alias(col))
        return df

    # —— 推广数据 CSV ——
    # PROMO_COL_MAP：列名 -> 统一长表列名。
    # 按**列名**选取而不是按位置，避免不同店铺导出的列顺序不一致导致取错列
    PROMO_COL_MAP = {
        "日期": "date",
        "SPU ID": "spu",
        "花费": "promotion_cost",  # 页面卡片「推广费」
        "总订单金额": "promotion_amount",  # 页面卡片「推广成交金额」
        # 「投产比」无需落库：ROI 由 transforms 按 推广成交金额/推广费 计算（= 投产比）
    }

    def _read_promo_one(self, path: Path, shop: str) -> pl.DataFrame | None:
        """读取单个推广 CSV，归一成 [shop, date, spu, promotion_cost, promotion_amount]。

        Args:
            path: 推广 CSV 文件路径。
            shop: 店铺名（由文件名推断），写入 shop 列。

        Returns:
            pl.DataFrame | None: 规整后的推广表；三种情况返回 None——
                所有编码都读不了、文件为空、或没有任何 PROMO_COL_MAP 里登记的列。

        Raises:
            无（读取异常被内部逐个编码吞掉，转为返回 None）。

        Example:
            >>> df = src._read_promo_one(Path("钻芯旗舰店_推广数据.csv"), "钻芯旗舰店")  # doctest: +SKIP
        """
        raw = None
        # 编码回退顺序：utf-8-sig（带 BOM 的导出）→ gbk（国内工具常见）→ utf-8。
        # 必须逐个试，否则 gbk 编码的中文列名会读成乱码，进而匹配不到列名
        for enc in ("utf-8-sig", "gbk", "utf-8"):
            try:
                # infer_schema_length=0 让所有列先按字符串读，避免数字列被推断成 Int 后 JOIN 失配
                raw = pl.read_csv(path, encoding=enc, infer_schema_length=0)
                break
            except Exception:
                raw = None
        if raw is None or raw.height == 0:
            return None
        keep = [c for c in self.PROMO_COL_MAP if c in raw.columns]
        if not keep:
            return None
        df = raw.select(keep).rename({c: self.PROMO_COL_MAP[c] for c in keep})
        # 与明细表同样的「剔除合计行 / 空行」处理
        df = df.filter(
            pl.col("spu").is_not_null()
            & (pl.col("spu").cast(pl.String) != "合计")
        )
        # 日期 YYYYMMDD -> YYYY-MM-DD；SPU 归整（与明细表一致，保证 JOIN 命中）
        df = df.with_columns(
            pl.lit(shop).alias("shop"),
            pl.col("date").cast(pl.String).map_elements(
                lambda d: f"{d[:4]}-{d[4:6]}-{d[6:8]}"
                # 只有纯 8 位数字才做转换，其它格式（已带横杠的）原样保留
                if (isinstance(d, str) and d.isdigit() and len(d) == 8)
                else (str(d) if d is not None else None),
                return_dtype=pl.String,
            ).alias("date"),
            pl.col("spu").cast(pl.String).map_elements(_norm_spu, return_dtype=pl.String).alias("spu"),
        )
        # 金额列同样去千分位/百分号后转 float
        for col in ("promotion_cost", "promotion_amount"):
            if col in df.columns:
                df = df.with_columns(
                    pl.col(col).cast(pl.String)
                    .str.replace_all(",", "")
                    .str.replace_all("%", "")
                    .cast(pl.Float64, strict=False)
                    .alias(col)
                )
        return df

    def _fetch_promotion(self) -> pl.DataFrame | None:
        """读取全部「*_推广数据_*.csv」，去重 + 白名单过滤后返回可直接 JOIN 的推广长表。

        Returns:
            pl.DataFrame | None: 列含 shop / date / spu / promotion_cost / promotion_amount 的表；
                目录不存在或没有任何可用推广文件时返回 None（调用方会跳过 JOIN）。

        Example:
            >>> df = src._fetch_promotion()   # doctest: +SKIP
        """
        table_dir = self.root / config.POP_TABLE_DIR_NAME
        if not table_dir.is_dir():
            return None
        frames = [
            self._read_promo_one(p, p.name.split("_")[0])
            # 同样跳过 ~$ 锁文件；按文件名排序保证「保留最后出现」的语义稳定
            for p in sorted(q for q in table_dir.glob("*推广数据*.csv") if not q.name.startswith("~$"))
        ]
        frames = [f for f in frames if f is not None]
        if not frames:
            return None
        df = pl.concat(frames, how="diagonal_relaxed")
        # 补齐列：某些店铺的 CSV 可能只有花费没有成交金额
        for col in ("shop", "date", "spu", "promotion_cost", "promotion_amount"):
            if col not in df.columns:
                df = df.with_columns(pl.lit(None).alias(col))
        # 去重：多份推广表日期区间重叠（如钻芯旗舰店两份），按 (shop,date,spu) 保留最后出现的
        df = df.unique(subset=["shop", "date", "spu"], keep="last")
        # 白名单过滤（与明细完全一致——两侧口径必须相同，否则 JOIN 后会出现单边数据）
        whitelist = load_spu_whitelist(self.root)
        if whitelist:
            allowed = {
                (s.strip(), _norm_spu(sp))
                for s, sps in whitelist.items()
                for sp in sps
            }
            df = df.filter(
                pl.struct(["shop", "spu"]).map_elements(
                    lambda s: (s["shop"].strip(), _norm_spu(s["spu"])) in allowed,
                    return_dtype=pl.Boolean,
                )
            )
        return df
